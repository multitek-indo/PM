"""Bounded XLSX reader. No database, formula evaluation, macros or network access."""
import io
import math
import re
import zipfile
from datetime import date, datetime, time
from pathlib import PurePosixPath

FIELDS = ['code', 'name', 'category', 'manufacturer', 'model', 'serial', 'site', 'line']
MAX_XLSX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 32
MAX_UNCOMPRESSED = 25 * 1024 * 1024
ALIASES = {
    'kode': 'code', 'kode mesin': 'code', 'nama': 'name', 'nama mesin': 'name',
    'jenis': 'category', 'jenis mesin': 'category', 'kategori': 'category',
    'merek': 'manufacturer', 'merk': 'manufacturer', 'serial number': 'serial',
    'nomor seri': 'serial', 'nomor serial': 'serial', 'pabrik': 'site',
    'lokasi': 'site', 'pabrik / lokasi': 'site', 'line produksi': 'line',
}

class ImportFileError(ValueError):
    """Safe, user-facing validation message. Never includes raw workbook XML."""

def check_archive(data):
    if not data or len(data) > MAX_XLSX_BYTES:
        raise ImportFileError('File Excel harus terisi dan maksimal 5 MiB.')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            items = archive.infolist()
            if len(items) > 1024 or sum(i.file_size for i in items) > MAX_UNCOMPRESSED:
                raise ImportFileError('Isi Excel terlalu besar setelah dibuka. Pecah menjadi beberapa file.')
            seen = set()
            for item in items:
                name = item.filename
                if name in seen or PurePosixPath(name).is_absolute() or '..' in PurePosixPath(name).parts:
                    raise ImportFileError('Struktur arsip Excel tidak valid.')
                seen.add(name)
                lower = name.lower()
                if item.flag_bits & 1 or item.file_size > 12*1024*1024:
                    raise ImportFileError('Excel terenkripsi atau salah satu bagian file terlalu besar.')
                if 'vbaproject' in lower or 'externallinks/' in lower or 'embeddings/' in lower:
                    raise ImportFileError('Macro, tautan workbook eksternal, dan objek tertanam tidak diizinkan.')
                if lower.endswith(('.xml', '.rels')):
                    content = archive.read(item)
                    if b'<!DOCTYPE' in content.upper() or b'<!ENTITY' in content.upper():
                        raise ImportFileError('Deklarasi XML eksternal tidak diizinkan.')
            if '[Content_Types].xml' not in seen or 'xl/workbook.xml' not in seen:
                raise ImportFileError('File bukan workbook XLSX yang valid.')
    except (zipfile.BadZipFile, RuntimeError, NotImplementedError, OSError) as exc:
        raise ImportFileError('File Excel rusak, terkunci, atau bukan format .xlsx.') from exc


def header_map(values):
    values = list(values)
    while values and values[-1] in (None, ''):
        values.pop()
    if len(values) > MAX_COLUMNS:
        raise ImportFileError('Terlalu banyak kolom; gunakan delapan kolom template.')
    names = []
    for value in values:
        key = re.sub(r'\s+', ' ', str(value or '').strip().lower().replace('_', ' '))
        key = ALIASES.get(key, key)
        if key not in FIELDS:
            raise ImportFileError('Header Excel tidak dikenal. Gunakan template: '+', '.join(FIELDS))
        names.append(key)
    if len(set(names)) != len(names):
        raise ImportFileError('Ada nama kolom berulang dalam header Excel.')
    if set(names) != set(FIELDS):
        raise ImportFileError('Header wajib lengkap: '+', '.join(FIELDS))
    return names


def cell_text(cell, field, row_number):
    value = cell.value
    where = f'Baris {row_number}, kolom {field}'
    if cell.data_type == 'f':
        raise ImportFileError(where+': rumus tidak diimpor. Gunakan Paste Special > Values di Excel.')
    if cell.data_type == 'e':
        raise ImportFileError(where+': sel berisi error Excel.')
    if value is None:
        return ''
    if isinstance(value, bool) or isinstance(value, (date, datetime, time)):
        raise ImportFileError(where+': gunakan teks, bukan tanggal atau nilai logika.')
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or (field in ('code', 'serial') and (value != int(value) or abs(value) >= 10**15)):
            raise ImportFileError(where+': kode/serial harus berupa teks agar identitas tidak berubah.')
        if value == int(value):
            text = str(int(value))
            fmt = cell.number_format or ''
            if field in ('code', 'serial') and value >= 0 and re.fullmatch(r'0{2,30}', fmt):
                text = text.zfill(len(fmt))
            return text
        return str(value)
    text = str(value).strip()
    if len(text) > 10000:
        raise ImportFileError(where+': teks terlalu panjang.')
    return text


def read_xlsx(data, sheet_name=''):
    check_archive(data)
    from openpyxl import load_workbook
    book = None
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        if len(book.sheetnames) > 30:
            raise ImportFileError('Maksimal 30 sheet dalam satu workbook.')
        visible = [s for s in book.worksheets if s.sheet_state == 'visible']
        if not visible:
            raise ImportFileError('Tidak ada sheet yang terlihat.')
        requested = sheet_name.strip()
        if requested:
            matches = [s for s in visible if s.title == requested]
            if not matches:
                raise ImportFileError('Nama sheet tidak ditemukan atau sheet tersembunyi. Periksa nama persisnya.')
            sheet = matches[0]
        else:
            sheet = next((s for s in visible if s.title.casefold() == 'mesin'), visible[0])
        if sheet.max_column and sheet.max_column > MAX_COLUMNS:
            raise ImportFileError('Sheet melebihi 32 kolom. Sisakan kolom template saja.')
        if sheet.max_row and sheet.max_row > MAX_ROWS+1:
            raise ImportFileError('Maksimal 5.000 baris data + 1 header. Hapus baris/format berlebih di Excel.')
        # Do not trust an underestimated dimension from a third-party exporter.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        try:
            header = next(iterator)
        except StopIteration:
            raise ImportFileError('Sheet kosong.')
        for cell in header:
            if cell.data_type in ('f', 'e'):
                raise ImportFileError('Header harus teks, bukan rumus/error.')
        names = header_map([c.value for c in header])
        rows = []
        for n, cells in enumerate(iterator, 2):
            if n > MAX_ROWS+1:
                raise ImportFileError('Maksimal 5.000 baris data per sheet.')
            if len(cells) > MAX_COLUMNS:
                raise ImportFileError(f'Baris {n}: terlalu banyak kolom.')
            if all(c.value in (None, '') for c in cells):
                continue
            if any(c.value not in (None, '') for c in cells[len(names):]):
                raise ImportFileError(f'Baris {n}: ada data di luar kolom header.')
            row = {}
            for i, key in enumerate(names):
                row[key] = cell_text(cells[i], key, n) if i < len(cells) else ''
            rows.append((n, row))
        if not rows:
            raise ImportFileError('Sheet tidak berisi data mesin. Isi mulai baris kedua.')
        return sheet.title, rows
    except ImportFileError:
        raise
    except Exception as exc:
        # Parsing third-party files may raise several library-specific exceptions.
        raise ImportFileError('Excel tidak dapat dibaca. Simpan ulang sebagai .xlsx tanpa password/macro.') from exc
    finally:
        if book is not None:
            book.close()


def template_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.comments import Comment
    from openpyxl.utils import get_column_letter
    book = Workbook()
    sheet = book.active
    sheet.title = 'Mesin'
    headings = ['Kode mesin (wajib; teks)', 'Nama mesin (wajib)', 'Jenis mesin', 'Merek', 'Model', 'Serial number (teks)', 'Pabrik/lokasi', 'Line produksi']
    widths = [23, 32, 21, 23, 23, 27, 25, 24]
    for i, field in enumerate(FIELDS, 1):
        c = sheet.cell(1, i, field)
        c.font = Font(bold=True, color='FFFFFF', size=11)
        c.fill = PatternFill('solid', fgColor='1264A3')
        c.alignment = Alignment(vertical='center')
        c.comment = Comment(headings[i-1], 'Multitek PM')
        sheet.column_dimensions[get_column_letter(i)].width = widths[i-1]
    for row in sheet.iter_rows(min_row=2, max_row=101, max_col=8):
        for c in row:
            c.number_format = '@'
            c.font = Font(color='135C9A', size=11)
            if c.row % 2 == 0:
                c.fill = PatternFill('solid', fgColor='EFF8FF')
    sheet.row_dimensions[1].height = 29
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:H101'
    sheet.sheet_view.showGridLines = False
    guide = book.create_sheet('Petunjuk')
    lines = [
        ('PANDUAN IMPOR MESIN', 'Multitek Smart Maintenance 0.2'),
        ('1. Isi sheet Mesin', 'Delapan header pada baris 1 tidak perlu diubah.'),
        ('2. Identitas', 'code dan serial sebaiknya berformat Text, terutama kode dengan nol di depan.'),
        ('3. Kolom wajib', 'code dan name harus terisi. Kolom lain boleh kosong.'),
        ('4. Batas', 'Maksimal 5.000 baris data dan ukuran file 5 MiB.'),
        ('5. Rumus', 'Gunakan nilai biasa. Rumus, macro, dan workbook eksternal ditolak.'),
        ('6. Keamanan data', 'Kode yang sudah ada membatalkan seluruh impor. Data lama tidak ditimpa.'),
        ('7. Jadwal', 'Impor hanya master mesin. Buat jadwal PM secara terpisah.'),
        ('8. Contoh', 'Sheet Contoh berisi data fiktif, tidak dipilih otomatis untuk impor.'),
    ]
    for row in lines: guide.append(row)
    guide.column_dimensions['A'].width = 29
    guide.column_dimensions['B'].width = 90
    guide.sheet_view.showGridLines = False
    for row in guide:
        guide.row_dimensions[row[0].row].height = 35
        for c in row:
            c.alignment = Alignment(vertical='center', wrap_text=True)
            c.font = Font(color='234867', size=11, bold=c.column == 1)
            if c.row == 1:
                c.fill = PatternFill('solid', fgColor='C9E8FF')
    example = book.create_sheet('Contoh')
    example.append(FIELDS)
    example.append(['CONTOH-001', 'Mesin jahit contoh', 'Sewing', 'Merek contoh', 'Model contoh', '0000123', 'Pabrik contoh', 'LINE-UJI'])
    for i, width in enumerate(widths, 1): example.column_dimensions[get_column_letter(i)].width = width
    for c in example[1]: c.fill = PatternFill('solid', fgColor='C9E8FF'); c.font = Font(bold=True, color='163D61')
    example.freeze_panes = 'A2'; example.sheet_view.showGridLines = False
    stream = io.BytesIO(); book.save(stream); book.close()
    return stream.getvalue()
