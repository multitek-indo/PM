"""Pure XLSX tests: no Django or database required."""
import io
import unittest
import zipfile
from datetime import date
from openpyxl import Workbook, load_workbook
from .tabular import FIELDS, ImportFileError, read_xlsx, template_xlsx, header_map, MAX_XLSX_BYTES

def workbook_bytes(rows=None, header=None, title='Mesin', edit=None):
    book=Workbook(); ws=book.active; ws.title=title
    ws.append(FIELDS if header is None else header)
    for row in rows or []: ws.append(row)
    if edit: edit(book, ws)
    out=io.BytesIO(); book.save(out); book.close(); return out.getvalue()

def row(code='EXAMPLE-001', name='Mesin contoh'):
    return [code,name,'Sewing','Merek contoh','Model uji','000123','Pabrik contoh','LINE-UJI']

class TabularTests(unittest.TestCase):
    def test_valid_workbook(self):
        sheet, rows=read_xlsx(workbook_bytes([row()]))
        self.assertEqual(sheet,'Mesin'); self.assertEqual(rows[0][1]['serial'],'000123')
    def test_indonesian_headers(self):
        h=['Kode mesin','Nama mesin','Jenis mesin','Merek','Model','Nomor seri','Pabrik','Line produksi']
        self.assertEqual(read_xlsx(workbook_bytes([row()],h))[1][0][1]['code'],'EXAMPLE-001')
    def test_header_case_and_whitespace(self):
        self.assertEqual(header_map([' CODE ','NAME','CATEGORY','MANUFACTURER','MODEL','SERIAL','SITE','LINE']),FIELDS)
    def test_missing_header(self):
        with self.assertRaises(ImportFileError): read_xlsx(workbook_bytes([row()],FIELDS[:7]))
    def test_unknown_header(self):
        with self.assertRaises(ImportFileError): read_xlsx(workbook_bytes([row()],FIELDS+['password']))
    def test_duplicate_header(self):
        with self.assertRaises(ImportFileError): read_xlsx(workbook_bytes([row()],FIELDS[:-1]+['code']))
    def test_reordered_header(self):
        sheet,rows=read_xlsx(workbook_bytes([row()[::-1]],FIELDS[::-1])); self.assertEqual(rows[0][1]['code'],'EXAMPLE-001')
    def test_empty_workbook(self):
        with self.assertRaises(ImportFileError): read_xlsx(workbook_bytes())
    def test_invalid_zip(self):
        with self.assertRaises(ImportFileError): read_xlsx(b'not-an-excel-file')
    def test_too_big(self):
        with self.assertRaises(ImportFileError): read_xlsx(b'X'*(MAX_XLSX_BYTES+1))
    def test_formula_rejected(self):
        with self.assertRaises(ImportFileError): read_xlsx(workbook_bytes([row('=1+1')]))
    def test_error_rejected(self):
        with self.assertRaises(ImportFileError): read_xlsx(workbook_bytes([row('#N/A')]))
    def test_dates_rejected(self):
        with self.assertRaises(ImportFileError): read_xlsx(workbook_bytes([row(date(2026,1,1))]))
    def test_boolean_rejected(self):
        with self.assertRaises(ImportFileError): read_xlsx(workbook_bytes([row(True)]))
    def test_large_numeric_code_rejected(self):
        with self.assertRaises(ImportFileError): read_xlsx(workbook_bytes([row(1234567890123456)]))
    def test_fractional_code_rejected(self):
        with self.assertRaises(ImportFileError): read_xlsx(workbook_bytes([row(12.5)]))
    def test_zero_padding(self):
        def edit(book,ws): ws['A2'].number_format='000000'
        self.assertEqual(read_xlsx(workbook_bytes([row(123)],edit=edit))[1][0][1]['code'],'000123')
    def test_numeric_no_decimal_suffix(self):
        self.assertEqual(read_xlsx(workbook_bytes([row(123.0)]))[1][0][1]['code'],'123')
    def test_blank_rows_skipped(self):
        rows=read_xlsx(workbook_bytes([row(),[],row('EXAMPLE-002')]))[1]
        self.assertEqual(len(rows),2); self.assertEqual(rows[1][0],4)
    def test_5001_rows_rejected(self):
        with self.assertRaises(ImportFileError): read_xlsx(workbook_bytes([row(str(i)) for i in range(5001)]))
    def test_extra_data_columns(self):
        with self.assertRaises(ImportFileError): read_xlsx(workbook_bytes([row()+['unexpected']]))
    def test_explicit_sheet(self):
        def edit(book,ws):
            other=book.create_sheet('Dipilih'); other.append(FIELDS); other.append(row('SECOND'))
        self.assertEqual(read_xlsx(workbook_bytes([row()],edit=edit),'Dipilih')[1][0][1]['code'],'SECOND')
    def test_missing_sheet(self):
        with self.assertRaises(ImportFileError): read_xlsx(workbook_bytes([row()]),'TidakAda')
    def test_hidden_sheet_rejected(self):
        def edit(book,ws):
            ws.sheet_state='hidden'; book.create_sheet('Terlihat')
        with self.assertRaises(ImportFileError): read_xlsx(workbook_bytes([row()],edit=edit),'Mesin')
    def test_macro_archive_rejected(self):
        data=io.BytesIO(workbook_bytes([row()]))
        with zipfile.ZipFile(data,'a') as z: z.writestr('xl/vbaProject.bin',b'dummy')
        with self.assertRaises(ImportFileError): read_xlsx(data.getvalue())
    def test_external_link_archive_rejected(self):
        data=io.BytesIO(workbook_bytes([row()]))
        with zipfile.ZipFile(data,'a') as z: z.writestr('xl/externalLinks/externalLink1.xml','<x/>')
        with self.assertRaises(ImportFileError): read_xlsx(data.getvalue())
    def test_dtd_rejected(self):
        data=io.BytesIO(workbook_bytes([row()]))
        with zipfile.ZipFile(data,'a') as z: z.writestr('extra.xml','<!DOCTYPE x [<!ENTITY a "x">]><x/>')
        with self.assertRaises(ImportFileError): read_xlsx(data.getvalue())
    def test_template_is_empty_and_valid(self):
        data=template_xlsx(); book=load_workbook(io.BytesIO(data))
        self.assertEqual(book.sheetnames,['Mesin','Petunjuk','Contoh'])
        self.assertEqual([c.value for c in book['Mesin'][1]],FIELDS)
        self.assertEqual(book['Mesin']['A2'].number_format,'@'); book.close()
        with self.assertRaises(ImportFileError): read_xlsx(data)
    def test_incorrect_dimension_is_not_trusted(self):
        original=io.BytesIO(workbook_bytes([row(),row('SECOND')]))
        out=io.BytesIO()
        with zipfile.ZipFile(original) as zin, zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as zout:
            for i in zin.infolist():
                raw=zin.read(i.filename)
                if i.filename=='xl/worksheets/sheet1.xml': raw=raw.replace(b'ref="A1:H3"',b'ref="A1:A1"')
                zout.writestr(i.filename,raw)
        self.assertEqual(len(read_xlsx(out.getvalue())[1]),2)
