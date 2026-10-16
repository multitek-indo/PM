"""Strict, bounded file reader. Atomic updates live only in inventory.services."""
import csv
import hashlib
import io
import re
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from django.core.exceptions import ValidationError

HEADERS = ['kode_sparepart', 'jenis', 'jumlah', 'referensi']
MAX_ROWS = 5000
MAX_XLSX = 5*1024*1024
MAX_CSV = 2*1024*1024


def parse_uploaded(upload):
    filename = Path(upload.name).name
    ext = Path(filename).suffix.lower()
    if ext not in ('.csv', '.xlsx'):
        raise ValidationError('Format harus .xlsx atau .csv. Tidak mendukung .xls/.xlsm.')
    data = upload.read((MAX_XLSX if ext=='.xlsx' else MAX_CSV)+1)
    if len(data) > (MAX_XLSX if ext=='.xlsx' else MAX_CSV):
        raise ValidationError('File terlalu besar (CSV 2 MiB; XLSX 5 MiB).')
    if not data:
        raise ValidationError('File kosong.')
    rows = read_xlsx(data) if ext == '.xlsx' else read_csv(data)
    return rows, hashlib.sha256(data).hexdigest(), filename


def read_csv(data):
    try:
        content = data.decode('utf-8-sig')
    except UnicodeDecodeError:
        raise ValidationError('CSV harus UTF-8.')
    if '\x00' in content:
        raise ValidationError('CSV mengandung byte yang tidak valid.')
    first = content.splitlines()[0] if content.splitlines() else ''
    delim = ';' if first.count(';') > first.count(',') else ','
    return normalize(csv.reader(io.StringIO(content, newline=''), delimiter=delim))


def read_xlsx(data):
    if not data.startswith(b'PK'):
        raise ValidationError('File XLSX tidak valid.')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries=archive.infolist()
            if len(entries)>200 or sum(e.file_size for e in entries)>24*1024*1024:
                raise ValidationError('Workbook terlalu besar setelah didekompresi.')
            if any(e.flag_bits & 1 or e.file_size > 18*1024*1024 for e in entries):
                raise ValidationError('Workbook terenkripsi atau melebihi batas.')
            if any('externalLink' in e.filename or 'vbaProject' in e.filename for e in entries):
                raise ValidationError('Link eksternal atau macro tidak diperbolehkan.')
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            ws=wb['Mutasi'] if 'Mutasi' in wb.sheetnames else next((s for s in wb.worksheets if s.sheet_state=='visible'),None)
            if ws is None:
                raise ValidationError('Tidak ada sheet yang terlihat.')
            if ws.max_row and ws.max_row>MAX_ROWS+2:
                raise ValidationError('Maksimal 5.000 baris mutasi.')
            def data_rows():
                for row in ws.iter_rows(min_row=1, max_col=5):
                    if any(c.data_type == 'f' for c in row):
                        raise ValidationError('Rumus Excel tidak diperbolehkan; Paste Special > Values terlebih dahulu.')
                    yield [c.value for c in row]
            return normalize(data_rows())
        finally:
            wb.close()
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValidationError('File Excel rusak atau tidak didukung.') from exc


def normalize(records):
    iterator=iter(records)
    head=next(iterator,None)
    if not head:
        raise ValidationError('File tidak mempunyai header.')
    h=[str(x or '').strip().lower() for x in head]
    if h[:4]!=HEADERS or any(str(x or '').strip() for x in h[4:]):
        raise ValidationError('Header harus: kode_sparepart, jenis, jumlah, referensi.')
    out=[]
    for index,row in enumerate(iterator,start=2):
        if index>MAX_ROWS+1:
            raise ValidationError('Maksimal 5.000 baris.')
        items=list(row)
        if len(items)>4 and any(str(x or '').strip() for x in items[4:]):
            raise ValidationError('Baris %d: kolom tambahan tidak diperbolehkan.' % index)
        items=(items+[None]*4)[:4]
        if not any(str(x or '').strip() for x in items):
            continue
        sku,kind,qty,note=items
        sku=str(sku or '').strip().upper()
        note=str(note or '').strip()
        kind=str(kind or '').strip().lower()
        if kind in ('masuk','barang masuk','receive'):kind='receive'
        elif kind in ('keluar','barang keluar','issue'):kind='issue'
        if not sku or len(sku)>64 or kind not in ('receive','issue') or not note or len(note)>600:
            raise ValidationError('Baris %d: kode, jenis, atau referensi tidak valid.' % index)
        if isinstance(qty,bool) or qty is None:
            raise ValidationError('Baris %d: jumlah harus bulat positif.' % index)
        try:
            num=Decimal(str(qty).strip())
        except (InvalidOperation,ValueError,TypeError):
            raise ValidationError('Baris %d: jumlah bukan angka.' % index)
        if not num.is_finite() or num<1 or num>999999999 or num!=num.to_integral_value():
            raise ValidationError('Baris %d: jumlah harus bilangan bulat 1-999999999.' % index)
        if isinstance(qty,str) and not re.fullmatch(r'[0-9]{1,9}', qty.strip()):
            raise ValidationError('Baris %d: gunakan angka bulat tanpa pemisah ribuan.' % index)
        if note[0] in '=+@':
            raise ValidationError('Baris %d: referensi tidak boleh berupa rumus.' % index)
        out.append({'sku':sku,'kind':kind,'quantity':int(num),'note':note})
    if not out:
        raise ValidationError('Tidak ada transaksi yang dapat diimpor.')
    return out
